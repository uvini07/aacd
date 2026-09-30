"""
Leitura da planilha bruta da AACD.

Na planilha original cada paciente ocupa DUAS linhas, com milhares de celulas
mescladas (3.839 na aba 60+). Ler com pandas direto devolve o dobro de linhas,
metade em branco. Este modulo resolve so isso: devolve uma linha por paciente,
com os valores exatamente como foram digitados.

Nenhuma conversao de tipo ou correcao acontece aqui — ver scripts/tratamento.py.
"""

import os
import re
from collections import OrderedDict

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PLANILHA = os.path.join(RAIZ, "base_quedas_AACD_anonimizada.xlsx")
SAIDAS = os.path.join(RAIZ, "saidas")
ABAS = ("60+", "60-")

# Marcadores de ausencia escritos a mao durante o evento.
# 'NA' foi usado para "teste nao realizado"; '.' aparece em suporte social e o
# significado ainda depende de confirmacao da AACD (ver pendencias).
AUSENTES_EXPLICITOS = {"NA", "N/A", "na", "n/a"}
AUSENTES_AMBIGUOS = {".", "-", "?", "k", "o"}


def normaliza_texto(valor):
    """Colapsa espacos repetidos e devolve None para vazio. Nao altera conteudo."""
    if valor is None:
        return None
    if isinstance(valor, str):
        texto = re.sub(r"\s+", " ", valor).strip()
        return texto or None
    return valor


def le_aba(nome_aba):
    """Devolve (cabecalhos, pacientes) com uma entrada por paciente.

    Cada paciente e um bloco de duas linhas. Onde a celula foi mesclada, o valor
    esta na primeira linha; onde nao foi, pode estar na segunda. As duas linhas
    nunca se contradizem (verificado nos 282 blocos), entao o primeiro valor
    nao-vazio do par e o valor do paciente.
    """
    planilha = openpyxl.load_workbook(PLANILHA, data_only=True)
    aba = planilha[nome_aba]
    cabecalhos = [
        re.sub(r"\s+", " ", str(aba.cell(1, coluna).value)).strip()
        for coluna in range(1, aba.max_column + 1)
    ]

    pacientes = []
    for linha in range(2, aba.max_row + 1, 2):
        registro = OrderedDict()
        for coluna, cabecalho in enumerate(cabecalhos, start=1):
            primeira = normaliza_texto(aba.cell(linha, coluna).value)
            segunda = (
                normaliza_texto(aba.cell(linha + 1, coluna).value)
                if linha + 1 <= aba.max_row
                else None
            )
            registro[cabecalho] = primeira if primeira is not None else segunda
        registro["_aba"] = nome_aba
        registro["_linha_excel"] = linha
        pacientes.append(registro)
    return cabecalhos, pacientes


def le_tudo():
    """Devolve {nome_aba: (cabecalhos, pacientes)} para as duas abas."""
    return {nome: le_aba(nome) for nome in ABAS}
